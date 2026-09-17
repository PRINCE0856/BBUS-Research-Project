"""Offline harness: fake TomTom responses so the whole resolve_places.py path
(search -> verdicts -> workbook) is exercised without network or quota.
Not part of the deliverable - a test scaffold only."""
import json
import os
import sys

os.environ["TOMTOM_API_KEY"] = "FAKE_KEY_FOR_OFFLINE_TEST"
import resolve_places as rp

BLR = (12.9716, 77.5946)


class FakeResp:
    def __init__(self, payload):
        self.status_code = 200
        self._p = payload

    def json(self):
        return self._p

    def raise_for_status(self):
        pass


def make(name, addr, typ, score, lat, lon, poi=False):
    r = {"type": typ, "score": score, "position": {"lat": lat, "lon": lon},
         "address": {"freeformAddress": addr}}
    if poi:
        r["poi"] = {"name": name}
    if typ == "Geography":
        r["entityType"] = "MunicipalitySubdivision"
    return r


def fake_get(url, params=None, headers=None, timeout=None):
    q = url.split("/search/2/search/")[1].rsplit(".json", 1)[0]
    from urllib.parse import unquote
    q = unquote(q).lower()
    lim = params.get("limit", 5)

    if "bilekahalli" in q:
        res = [make("Bilekahalli", "Bilekahalli, Bengaluru 560076", "Geography", 4.9, 12.8905, 77.6015),
               make("Bilekahalli Bus Stop", "Bannerghatta Rd, Bengaluru", "POI", 3.1, 12.8899, 77.6002, poi=True)]
    elif "konappana agrahara" in q:
        res = [make("Konappana Agrahara", "Konappana Agrahara, Bengaluru 560100", "Geography", 4.7, 12.8447, 77.6631)]
    elif "electronic city" in q:
        res = [make("Electronic City", "Electronic City, Bengaluru 560100", "Geography", 4.8, 12.8452, 77.6601)]
    elif "no such" in q or "chalavadipalya" in q:
        res = []
    elif "challaghatta" in q:
        # deliberately far away, to exercise the implausible branch
        res = [make("Challaghatta", "Challaghatta, Bengaluru 560074", "Geography", 4.2, 12.9080, 77.4380)]
    elif "infosys" in q:
        res = [make("Infosys Limited", "Electronics City Phase 1, Bengaluru", "POI", 4.4, 12.8460, 77.6620, poi=True),
               make("Infosys", "Hebbal, Bengaluru", "POI", 3.9, 13.0350, 77.5970, poi=True)]
    elif "besant" in q:
        res = [make("Besant Technologies", "BTM Layout, Bengaluru", "POI", 4.1, 12.9160, 77.6100, poi=True)]
    elif "punjab" in q:
        res = [make("Punjab & Sind Bank", "Jayanagar, Bengaluru", "POI", 4.0, 12.9300, 77.5830, poi=True)]
    elif "bescom" in q or "electricity" in q:
        res = [make("BESCOM", "Electronic City, Bengaluru", "POI", 3.8, 12.8440, 77.6600, poi=True)]
    elif "tilak nagar" in q or "thilaknagar" in q or "tilakpura" in q:
        res = [make("Tilak Nagar", "Tilak Nagar, Bengaluru 560041", "Geography", 3.6, 12.9260, 77.5860)]
    elif "kottigepalya" in q:
        res = [make("Kottigepalya", "Kottigepalya, Bengaluru 560091", "Geography", 4.6, 12.9760, 77.4870)]
    elif "kodipalya" in q or "kottanur" in q or "kottapalya" in q:
        res = [make("Kodipalya", "Kodipalya, Bengaluru", "Geography", 3.3, 12.8990, 77.4460)]
    elif "konanakunte" in q:
        res = [make("Konanakunte", "Konanakunte, Bengaluru 560062", "Geography", 4.3, 12.8840, 77.5620)]
    else:
        # generic endpoint geocode (Whitefield, Sarjapur, BTM, etc.)
        res = [make(q[:30], f"{q[:30]}, Bengaluru", "Geography", 4.5, 12.95, 77.65)]
    return FakeResp({"results": res[:lim]})


rp.requests.get = fake_get
sys.argv = ["resolve_places.py", "--sleep", "0"]
rp.main()

# ---- verify the workbook actually opened and looks right ----
from openpyxl import load_workbook
wb = load_workbook(rp.Path(__file__).resolve().parent / "Place_Candidates.xlsx")
print("\n\n==== WORKBOOK CHECK ====")
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"  {name}: {ws.max_row} rows x {ws.max_column} cols")
ws = wb["Candidates"]
print("\nheader:", [c.value for c in ws[1]])
vcol = [c.value for c in ws[1]].index("Verdict")
pcol = [c.value for c in ws[1]].index("Ready-to-paste PLACES line")
verdicts = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[vcol]:
        v = str(r[vcol]).split(" - ")[0]
        verdicts[v] = verdicts.get(v, 0) + 1
print("verdict spread:", verdicts)
# no paste line may ever accompany a non-strong verdict
bad = [(r[0], r[vcol]) for r in ws.iter_rows(min_row=2, values_only=True)
       if r[pcol] and not str(r[vcol] or "").startswith("strong")]
print("paste lines on non-strong verdicts (must be empty):", bad)
