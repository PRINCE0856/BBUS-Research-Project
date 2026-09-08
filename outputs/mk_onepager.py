from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEAL = RGBColor(0x00, 0x6D, 0x77); DARK = RGBColor(0x1A,0x1A,0x1A); GREY = RGBColor(0x55,0x55,0x55)
BODY = 9.0

doc = Document()
s = doc.sections[0]
s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.2); s.left_margin = Cm(1.6); s.right_margin = Cm(1.6)
n = doc.styles['Normal']
n.font.name = 'Calibri'; n.font.size = Pt(BODY); n.font.color.rgb = DARK
n.paragraph_format.space_after = Pt(3); n.paragraph_format.line_spacing = 1.0

def para(text, size=BODY, italic=False, after=3, color=DARK):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic; r.font.color.rgb = color
    return p

def head(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper()); r.font.size = Pt(8.5); r.bold = True; r.font.color.rgb = TEAL
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'),'single'); bt.set(qn('w:sz'),'4'); bt.set(qn('w:space'),'2'); bt.set(qn('w:color'),'006D77')
    pb.append(bt); pPr.append(pb)
    return p

def bullet(lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.25)
    if lead:
        r = p.add_run(lead); r.bold = True; r.font.size = Pt(BODY)
    r = p.add_run(text); r.font.size = Pt(BODY)
    return p

# Title
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
r = p.add_run('#BBUS  |  Quantifying the full value of India’s bus systems')
r.font.size = Pt(14.5); r.bold = True; r.font.color.rgb = TEAL
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
r = p.add_run('Efficiency and Bus Fleet Management: strategic communications and quantifying GHG and other '
              'co-benefits  ·  CEEW  ·  One-page explainer, September 2026')
r.font.size = Pt(8); r.italic = True; r.font.color.rgb = GREY

head('The project in one sentence')
para('India is about to spend tens of thousands of crores on city buses, yet still justifies that spending almost '
     'entirely on fares and ridership. #BBUS builds the missing evidence — one integrated model that puts a rupee value '
     'on what a bus actually delivers to a city: better health, higher incomes and greater access for women and the '
     'poor, lower emissions, and mobility that survives a flood or a heatwave — so governments can defend bus budgets '
     'and unlock climate and private finance.')

head('Why this matters now')
bullet('More than 230 million Indians ', 'depend on buses every day. Most are captive users — women, the elderly, '
       'children, persons with disabilities and low-income households — with no realistic alternative.')
bullet('Over 60 per cent ', 'of India’s bus fleet sits in just nine metro cities. Elsewhere supply is thin, service '
       'inconsistent, investment reactive.')
bullet('PM e-Bus Sewa ', 'commits about 10,000 electric buses across 169 cities, an estimated ₹57,613 crore with '
       '₹20,000 crore of central support. It is the largest urban bus intervention since JNNURM, and the evidence it '
       'generates will shape the second phase already being planned.')
bullet('The gap: ', 'appraisal today counts only farebox revenue, operating cost and in-vehicle time saved. A scoping '
       'review of 27 studies found exactly one India-specific co-benefit study. Conventional cost-benefit analysis '
       'therefore understates the social return to bus investment, systematically.')

head('What we are building: the 4-in-1 co-benefits model')
rows = [
    ('Model', 'What it answers', 'Headline metrics'),
    ('M1 Socio-economic', 'Who gains, and how much, in money and opportunity — especially women, low-income and vulnerable road users?',
     'Household transport spend saved; female labour force participation; time saved; jobs reachable; decongestion and agglomeration value'),
    ('M2 Health', 'What is the health value of moving people out of two-wheelers and cars and onto buses?',
     'DALYs averted; walking gained; PM2.5 exposure avoided; road injuries avoided; commute stress reduced; treatment costs saved'),
    ('M3 GHG', 'How much carbon does bus deployment actually avoid under realistic grid and fuel scenarios?',
     'Vehicle-km avoided; tonnes CO₂e per year; scenario range by grid mix and charging pattern'),
    ('M4 Resilience', 'What is a bus network worth when a city floods, overheats or shuts down?',
     'Economic downtime avoided; evacuation time saved; infrastructure damage and restart cost by mode'),
]
t = doc.add_table(rows=len(rows), cols=3); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Cm(2.7), Cm(6.0), Cm(8.5)]
for i, row in enumerate(rows):
    for j, txt in enumerate(row):
        cell = t.cell(i, j); cell.width = widths[j]
        cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(0.5); cp.paragraph_format.space_before = Pt(0.5)
        r = cp.add_run(txt); r.font.size = Pt(7.8)
        if i == 0: r.bold = True; r.font.color.rgb = TEAL
        elif j == 0: r.bold = True

head('How we generate the evidence')
bullet('Group A, buses already running: ', 'Bengaluru, Jaipur, Bhubaneswar. No pre-bus baseline exists, so we use '
       'cross-sectional regression with rich controls, and a 500-metre walkshed of a bus stop as the access variable.')
bullet('Group B, buses arriving: ', 'Jodhpur, Haridwar, Gaya, Bilaspur, Tirupati, Puducherry, Gandhinagar (final list '
       'with MoHUA). The staggered PM e-Bus Sewa rollout is a natural experiment — we survey before deployment and '
       'again once service stabilises, giving a difference-in-differences estimate of what the buses caused.')
bullet('Data: ', 'a structured household and commuter survey, preceded by in-depth interviews and focus groups; '
       'passive mobile GPS trip data giving origin-destination, trip length, speed and route traces at city scale; GIS '
       'route, stop and hazard layers; and secondary data from the Census, PLFS, Economic Census and air-quality networks.')
bullet('Equity is in the specification: ', 'a four-part vulnerability index (income, vehicle ownership, disability, '
       'elderly status) enters every regression as an interaction, so “who benefits most” is a headline coefficient, '
       'not a footnote.')

head('What the project produces')
bullet('A Pre-Investment Appraisal Tool. ', 'Open-access and online. A city official enters fleet size, population, '
       'modal share and grid emission factor, and gets back projected GHG savings, health co-benefits, an SROI ratio, '
       'the climate finance instruments the city qualifies for, and an MRV dashboard aligned to India’s NDC and the '
       'SDGs. It is formatted to meet GCF, ADB and World Bank project preparation requirements — precisely where '
       'smaller cities get stuck.')
bullet('A Healthy Cities SROI framework. ', 'The same dataset read three ways: for government (societal return per '
       'rupee of public spend, including what the ₹24/km PM e-Bus Sewa subsidy is actually buying); for climate and '
       'development finance (structured to GCF and ADB results frameworks); and for private capital (co-benefits sized '
       'as a payment security mechanism that de-risks private operators).')

head('The line to use in the room')
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(1); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run('“Today a city can tell you exactly what a bus costs. It cannot tell you what a bus is worth. '
              '#BBUS builds the number that closes that gap — and makes it bankable.”')
r.font.size = Pt(9.5); r.italic = True; r.bold = True; r.font.color.rgb = TEAL

doc.save('BBUS_Project_Explainer_1pager.docx')
print('saved OK')
